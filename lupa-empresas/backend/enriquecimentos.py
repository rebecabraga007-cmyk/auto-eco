# -*- coding: utf-8 -*-
"""Os enriquecimentos que ja foram feitos -- a lista que sobrevive ao navegador.

O QUE ISTO CONSERTA
-------------------
Hoje a planilha enriquecida existe apenas na memoria da ABA. Quem fecha o
navegador antes de clicar em "Exportar XLSX" perde tudo -- e "tudo" aqui e
trabalho ja PAGO, consulta por consulta. Um lote de mil linhas leva minutos;
esperar de olho na tela e a unica forma de nao perder.

A Datastone resolve isso com uma tela de "Enriquecimentos" (analisada em
17/set/2026): cada envio vira uma linha com nome, data, quantidade, status e
o arquivo para baixar, e o arquivo tambem chega por e-mail. Copiar isso e
copiar a parte que importa -- o resultado deixa de depender da aba ficar
aberta.

DIFERENCA DELES PARA NOS, DE PROPOSITO
--------------------------------------
La o processamento e do servidor: a pessoa envia e volta depois. Aqui o lote
ja roda em ondas no navegador, com barra de progresso, e isso FUNCIONA -- foi
o que resolveu o corte de 100s da Cloudflare. Entao nao troco o motor: quando
o ultimo lote fecha, a tela manda o resultado para ca e ele vira arquivo em
disco. O ganho e o mesmo (o resultado persiste), sem refazer o que ja esta de
pe.

O ARQUIVO NAO E REGERAVEL. Guardo o XLSX pronto, e nao as linhas para montar
depois: as linhas vieram de consultas pagas num instante especifico, e
"gerar de novo" seria cobrar de novo.
"""
import json
import os
import re
import sqlite3
import time
import uuid

_AQUI = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(os.environ.get("CAPIBLU_DATA_DIR", _AQUI), "enriquecimentos.db")
ARQUIVOS = os.path.join(os.environ.get("CAPIBLU_DATA_DIR", _AQUI),
                        "enriquecimentos_arquivos")

DDL = """
CREATE TABLE IF NOT EXISTS enriquecimento (
  id           TEXT PRIMARY KEY,
  usuario      TEXT,
  nome         TEXT,          -- o que a pessoa chamou
  origem       TEXT,          -- nome do arquivo que ela subiu
  status       TEXT DEFAULT 'concluido',
  linhas       INTEGER,       -- linhas da planilha
  enriquecidas INTEGER,       -- quantas voltaram com algo
  com_decisor  INTEGER,       -- quantas ficaram com pelo menos um decisor
  colunas      TEXT,          -- JSON [{key,label}] -- o cabecalho entregue
  parametros   TEXT,          -- JSON das quantidades escolhidas
  arquivo      TEXT,          -- nome do .xlsx no disco
  bytes        INTEGER,
  email_para   TEXT,
  email_em     INTEGER,
  email_erro   TEXT,
  criado_em    INTEGER
);
CREATE INDEX IF NOT EXISTS ix_enr_user ON enriquecimento(usuario, criado_em DESC);
"""


def _con():
    """Conexao que FECHA. `with sqlite3.connect()` nao fecha -- so faz commit.

    Ja custou caro neste projeto: 1023 descritores vazados derrubaram o
    servico inteiro com 502 e nenhum erro no log.
    """
    os.makedirs(os.path.dirname(DB) or ".", exist_ok=True)
    c = sqlite3.connect(DB, timeout=10)
    c.row_factory = sqlite3.Row
    return c


def _init():
    c = _con()
    try:
        c.executescript(DDL)
        c.commit()
    finally:
        c.close()
    os.makedirs(ARQUIVOS, exist_ok=True)


_init()


def salvar(usuario: str, nome: str, origem: str, colunas: list,
           conteudo: bytes, linhas: int = 0, enriquecidas: int = 0,
           com_decisor: int = 0, parametros: dict = None) -> dict:
    """Grava o XLSX em disco e registra a linha da lista."""
    eid = uuid.uuid4().hex[:12]
    seguro = re.sub(r"[^A-Za-z0-9._-]+", "-", (nome or "enriquecimento"))[:60]
    arq = "%s-%s.xlsx" % (eid, seguro or "enriquecimento")
    os.makedirs(ARQUIVOS, exist_ok=True)
    caminho = os.path.join(ARQUIVOS, arq)
    tmp = caminho + ".parcial"
    # Temporario + rename: renomear e atomico, entao ou o arquivo esta
    # inteiro ou nao existe. Meio arquivo na lista seria pior que nenhum --
    # a pessoa baixaria um XLSX corrompido achando que perdeu a consulta.
    with open(tmp, "wb") as f:
        f.write(conteudo)
    os.replace(tmp, caminho)

    c = _con()
    try:
        c.execute(
            "INSERT INTO enriquecimento (id, usuario, nome, origem, status,"
            " linhas, enriquecidas, com_decisor, colunas, parametros, arquivo,"
            " bytes, criado_em) VALUES (?,?,?,?,'concluido',?,?,?,?,?,?,?,?)",
            (eid, (usuario or "").strip().lower(), (nome or "").strip()[:120],
             (origem or "")[:160], int(linhas or 0), int(enriquecidas or 0),
             int(com_decisor or 0),
             json.dumps(colunas or [], ensure_ascii=False),
             json.dumps(parametros or {}, ensure_ascii=False),
             arq, len(conteudo), int(time.time())))
        c.commit()
    finally:
        c.close()
    return {"status": "ok", "id": eid, "arquivo": arq, "bytes": len(conteudo)}


def _linha(r: sqlite3.Row) -> dict:
    return {
        "id": r["id"], "nome": r["nome"] or "(sem nome)",
        "usuario": r["usuario"], "origem": r["origem"],
        "status": r["status"], "linhas": r["linhas"],
        "enriquecidas": r["enriquecidas"], "com_decisor": r["com_decisor"],
        "colunas": json.loads(r["colunas"] or "[]"),
        "parametros": json.loads(r["parametros"] or "{}"),
        "bytes": r["bytes"], "criado_em": r["criado_em"],
        "email_para": r["email_para"], "email_em": r["email_em"],
        "email_erro": r["email_erro"],
    }


def listar(usuario: str, admin: bool = False, limite: int = 60) -> list:
    """A lista de quem pediu. Admin ve de todo mundo.

    Sem admin, o filtro por usuario NAO e conveniencia de tela: a planilha
    enriquecida tem CPF e telefone de pessoa fisica, e a lista de outro
    operador nao e assunto de quem esta olhando.
    """
    c = _con()
    try:
        if admin:
            # `rowid` DESEMPATA. `criado_em` e em segundos, e um complemento
            # rapido termina no MESMO segundo da lista que o originou --
            # nesse caso a ordem ficava indefinida e a mais nova podia
            # aparecer embaixo. Quem clica em "completar" na primeira linha
            # pegaria a lista errada.
            rs = c.execute("SELECT * FROM enriquecimento"
                           " ORDER BY criado_em DESC, rowid DESC"
                           " LIMIT ?", (int(limite),)).fetchall()
        else:
            rs = c.execute("SELECT * FROM enriquecimento WHERE usuario = ?"
                           " ORDER BY criado_em DESC, rowid DESC LIMIT ?",
                           ((usuario or "").strip().lower(), int(limite))).fetchall()
        return [_linha(r) for r in rs]
    finally:
        c.close()


def pegar(eid: str, usuario: str, admin: bool = False) -> dict:
    """Um enriquecimento, com o caminho do arquivo. None se nao for dono."""
    if not re.fullmatch(r"[0-9a-f]{6,40}", str(eid or "")):
        return None            # id vem do cliente: nao vira caminho de arquivo
    c = _con()
    try:
        r = c.execute("SELECT * FROM enriquecimento WHERE id = ?", (eid,)).fetchone()
    finally:
        c.close()
    if not r:
        return None
    if not admin and (r["usuario"] or "") != (usuario or "").strip().lower():
        return None
    d = _linha(r)
    d["caminho"] = os.path.join(ARQUIVOS, r["arquivo"] or "")
    d["arquivo"] = r["arquivo"]
    return d


def marcar_email(eid: str, para: str, erro: str = "") -> None:
    """Registra a tentativa -- inclusive quando falhou.

    Guardar so o sucesso faz a lista mentir por omissao: quem nao recebeu
    veria o mesmo "—" de quem nunca pediu, e concluiria que esqueceu de
    clicar.
    """
    c = _con()
    try:
        c.execute("UPDATE enriquecimento SET email_para = ?, email_em = ?,"
                  " email_erro = ? WHERE id = ?",
                  (para, int(time.time()), (erro or "")[:200], eid))
        c.commit()
    finally:
        c.close()


def apagar(eid: str, usuario: str, admin: bool = False) -> bool:
    """Tira da lista e do disco. So o dono (ou admin)."""
    d = pegar(eid, usuario, admin)
    if not d:
        return False
    try:
        os.remove(d["caminho"])
    except OSError:
        pass
    c = _con()
    try:
        c.execute("DELETE FROM enriquecimento WHERE id = ?", (eid,))
        c.commit()
    finally:
        c.close()
    return True


def ler_linhas(eid: str, usuario: str = "", admin: bool = False) -> tuple:
    """As linhas de volta, lidas do PROPRIO XLSX que foi entregue.

    POR QUE DO ARQUIVO, e nao de uma copia guardada a parte: o XLSX e o que a
    pessoa recebeu e o que ela vai comparar. Se eu guardasse as linhas
    separadas e as duas divergissem -- por um ajuste de formatacao, por uma
    coluna renomeada -- o complemento sairia diferente do original e ninguem
    saberia dizer qual dos dois esta certo.

    O cabecalho do arquivo tem os ROTULOS ("Decisor 1 Telefone 1"), nao as
    chaves internas (`de_dec1_tel1`). A traducao de volta sai de `colunas`,
    que foi gravado junto no mesmo momento -- por isso ele existe.

    Devolve (registro, linhas). `registro` e None quando nao e da pessoa.
    """
    d = pegar(eid, usuario, admin)
    if not d or not os.path.exists(d.get("caminho") or ""):
        return None, []
    from openpyxl import load_workbook

    wb = load_workbook(d["caminho"], read_only=True, data_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            cabecalho = next(it)
        except StopIteration:
            return d, []
        de_volta = {c.get("label"): c.get("key") for c in (d["colunas"] or [])}
        chaves = [de_volta.get(h, h) for h in cabecalho]
        linhas = []
        for r in it:
            # Celula vazia vem como None, e None atravessa o codigo todo ate
            # virar a string "None" dentro de uma planilha de cliente.
            linhas.append({k: ("" if v is None else v)
                           for k, v in zip(chaves, r) if k})
    finally:
        wb.close()
    return d, linhas
